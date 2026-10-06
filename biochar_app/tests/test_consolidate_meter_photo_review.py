from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image

from biochar_app.scripts.management.consolidate_meter_photo_review import consolidate


def test_consolidation_updates_exact_sha_and_appends_new_photo(tmp_path: Path) -> None:
    originals = tmp_path / "originals"
    originals.mkdir()
    Image.new("RGB", (640, 480), "white").save(originals / "new.jpg")

    main = Workbook()
    sheet = main.active
    sheet.title = "Meter Review"
    sheet.append([
        "Photo Date/Time", "Photo", "Manual Reading", "Status", "Notes",
        "Filename", "SHA-256", "Crop Method", "Duplicate Count",
    ])
    sheet.append(["2026-01-01 00:00:00", "", "", "uncertain", "", "old.jpg", "a" * 64, "", 1])
    main_path = tmp_path / "main.xlsx"
    main.save(main_path)

    supplemental = Workbook()
    review = supplemental.active
    review.title = "Review"
    for _ in range(5):
        review.append([])
    review.append([
        "Photo", "Photo date/time", "Manual reading", "Status", "Notes",
        "Filename", "Timestamp confidence", "Timestamp source", "SHA-256",
        "Crop method", "Review check",
    ])
    review.append(["", "2026-01-01 00:00:00", "123456", "readable", "checked", "old.jpg", "high", "datetime_original", "a" * 64, "", "Complete"])
    review.append(["", "2026-01-02 00:00:00", "123457", "readable", "", "new.jpg", "high", "datetime_original", "b" * 64, "", "Complete"])
    supplemental_path = tmp_path / "supplemental.xlsx"
    supplemental.save(supplemental_path)

    output = tmp_path / "output.xlsx"
    result = consolidate(main_path, supplemental_path, originals, output)

    assert result == {
        "supplemental_rows": 2,
        "updated_rows": 1,
        "appended_rows": 1,
        "final_data_rows": 2,
    }
    workbook = load_workbook(output, data_only=True)
    output_sheet = workbook["Meter Review"]
    assert output_sheet["C2"].value == "123456"
    assert output_sheet["D2"].value == "readable"
    assert output_sheet["C3"].value == "123457"
    assert len(output_sheet._images) == 1
