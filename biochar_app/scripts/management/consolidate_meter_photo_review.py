#!/usr/bin/env python3
"""Merge a completed supplemental meter-photo review into the main workbook."""

from __future__ import annotations

import argparse
import os
import re
import tempfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill

from biochar_app.config.paths import IRRIGATION_DIR
from biochar_app.scripts.management.build_meter_review_workbook import (
    ROW_HEIGHT_POINTS,
    THUMBNAIL_HEIGHT_PX,
    THUMBNAIL_WIDTH_PX,
    fit_thumbnail,
    register_heic_support,
    save_workbook_crop,
)

PHOTOS_ROOT = IRRIGATION_DIR / "photos"
DEFAULT_MAIN = PHOTOS_ROOT / "meter_photo_review.xlsx"
DEFAULT_SUPPLEMENTAL = PHOTOS_ROOT / "meter_photo_unresolved_review.xlsx"
DEFAULT_ORIGINALS = PHOTOS_ROOT / "originals"


def clean(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none", "<na>"}:
        return ""
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def headers(sheet, row: int) -> dict[str, int]:
    return {clean(c.value): c.column for c in sheet[row] if clean(c.value)}


def copy_row_style(sheet, source_row: int, target_row: int) -> None:
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.border = copy(source.border)
        target.fill = copy(source.fill)
        target.font = copy(source.font)
        target.protection = copy(source.protection)


def extend_data_validations(sheet, new_last_row: int) -> None:
    for validation in sheet.data_validations.dataValidation:
        ranges = [str(r) for r in validation.ranges.ranges]
        if any(r.startswith("C2:C") for r in ranges):
            validation.sqref = f"C2:C{new_last_row}"
        elif any(r.startswith("D2:D") for r in ranges):
            validation.sqref = f"D2:D{new_last_row}"


def atomic_save(workbook, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        dir=output_path.parent,
        prefix=f".{output_path.stem}-",
        suffix=output_path.suffix,
        delete=False,
    ) as handle:
        temporary_path = Path(handle.name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def consolidate(
    main_workbook: Path,
    supplemental_workbook: Path,
    originals_dir: Path,
    output_workbook: Path,
) -> dict[str, int]:
    workbook = load_workbook(main_workbook)
    supplemental = load_workbook(supplemental_workbook, data_only=True, read_only=True)
    review = workbook["Meter Review"]
    supplemental_review = supplemental["Review"]
    main_headers = headers(review, 1)
    supplemental_headers = headers(supplemental_review, 6)

    required_main = {
        "Photo Date/Time", "Photo", "Manual Reading", "Status", "Notes",
        "Filename", "SHA-256", "Crop Method", "Duplicate Count",
    }
    required_supplemental = {
        "Photo date/time", "Manual reading", "Status", "Notes",
        "Filename", "SHA-256",
    }
    if missing := required_main.difference(main_headers):
        raise KeyError(f"Main workbook is missing columns: {sorted(missing)}")
    if missing := required_supplemental.difference(supplemental_headers):
        raise KeyError(f"Supplemental workbook is missing columns: {sorted(missing)}")

    by_sha: dict[str, int] = {}
    for row in range(2, review.max_row + 1):
        sha = clean(review.cell(row, main_headers["SHA-256"]).value).lower()
        if sha:
            by_sha[sha] = row

    supplemental_rows: list[dict[str, object]] = []
    supplemental_names = [
        clean(cell.value) for cell in supplemental_review[6]
    ]
    for values in supplemental_review.iter_rows(min_row=7, values_only=True):
        record = dict(zip(supplemental_names, values))
        if clean(record.get("SHA-256")):
            supplemental_rows.append(record)

    old_last_row = review.max_row
    source_style_row = old_last_row
    crop_dir = Path(tempfile.mkdtemp(prefix="meter-review-crops-"))
    heic_available = register_heic_support()
    updated = 0
    appended = 0

    for record in supplemental_rows:
        sha = clean(record["SHA-256"]).lower()
        existing_row = by_sha.get(sha)
        if existing_row is not None:
            review.cell(existing_row, main_headers["Manual Reading"]).value = clean(record["Manual reading"])
            review.cell(existing_row, main_headers["Status"]).value = clean(record["Status"])
            review.cell(existing_row, main_headers["Notes"]).value = clean(record["Notes"])
            updated += 1
            continue

        target_row = review.max_row + 1
        copy_row_style(review, source_style_row, target_row)
        review.row_dimensions[target_row].height = ROW_HEIGHT_POINTS
        filename = clean(record["Filename"])
        image_path = originals_dir / filename
        values = {
            "Photo Date/Time": record["Photo date/time"], "Photo": "",
            "Manual Reading": clean(record["Manual reading"]),
            "Status": clean(record["Status"]), "Notes": clean(record["Notes"]),
            "Filename": filename, "SHA-256": clean(record["SHA-256"]),
            "Crop Method": "", "Duplicate Count": 1,
        }
        for name, value in values.items():
            review.cell(target_row, main_headers[name]).value = value
        filename_cell = review.cell(target_row, main_headers["Filename"])
        filename_cell.hyperlink = image_path.resolve().as_uri()
        filename_cell.style = "Hyperlink"
        review.cell(target_row, main_headers["Manual Reading"]).number_format = "@"

        if not image_path.exists():
            raise FileNotFoundError(f"Original photo not found: {image_path}")
        crop_path, method = save_workbook_crop(image_path, crop_dir, heic_available)
        review.cell(target_row, main_headers["Crop Method"]).value = method
        width, height = fit_thumbnail(crop_path, THUMBNAIL_WIDTH_PX, THUMBNAIL_HEIGHT_PX)
        image = XLImage(str(crop_path))
        image.width, image.height, image.anchor = width, height, f"B{target_row}"
        review.add_image(image)
        by_sha[sha] = target_row
        appended += 1

    new_last_row = review.max_row
    extend_data_validations(review, new_last_row)
    review.auto_filter.ref = f"A1:I{new_last_row}"
    if new_last_row > old_last_row:
        new_range = f"A{old_last_row + 1}:I{new_last_row}"
        for formula, color in (
            (f'$C{old_last_row + 1}=""', "FFF2CC"),
            (f'$D{old_last_row + 1}="uncertain"', "FCE4D6"),
            (f'$D{old_last_row + 1}="readable"', "E2F0D9"),
        ):
            review.conditional_formatting.add(
                new_range,
                FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=color)),
            )
    atomic_save(workbook, output_workbook)
    return {
        "supplemental_rows": len(supplemental_rows), "updated_rows": updated,
        "appended_rows": appended, "final_data_rows": new_last_row - 1,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--main-workbook", type=Path, default=DEFAULT_MAIN)
    parser.add_argument("--supplemental-workbook", type=Path, default=DEFAULT_SUPPLEMENTAL)
    parser.add_argument("--originals-dir", type=Path, default=DEFAULT_ORIGINALS)
    parser.add_argument("--output-workbook", type=Path, default=DEFAULT_MAIN)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    result = consolidate(args.main_workbook, args.supplemental_workbook, args.originals_dir, args.output_workbook)
    for key, value in result.items():
        print(f"{key}: {value}")
    print(f"output: {args.output_workbook}")


if __name__ == "__main__":
    main()
