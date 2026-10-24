"""Tests for field-biomass and supplemental NIR ETL builders."""

from __future__ import annotations

import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from biochar_app.config.paths import WARD_MASTER_NIR_CSV
from biochar_app.scripts.lab.build_field_biomass_from_master import (
    BiomassDateWarning,
    build_field_biomass,
    extract_biomass_sheet,
)
from biochar_app.scripts.lab.update_ward_master_nir import (
    _read_supplemental_nir_csv,
)
from biochar_app.scripts.tables.tables_nir import build_nir_set1_table


LOCATIONS = [f"S{strip}{position}" for strip in range(1, 5) for position in "TMB"]


def _add_sheet(
    workbook: Workbook, name: str, year: int, old_layout: bool = False
) -> None:
    worksheet = workbook.create_sheet(name)
    if old_layout:
        worksheet.append(
            ["6-12-23 Biomass in grams", None, None, None, "7/25/23", None]
        )
        worksheet.append(["Location", "Wet", "24 hours", "48 hours", "Wet", "Dry"])
        for index, location in enumerate(LOCATIONS, start=1):
            worksheet.append(
                [location, 100 + index, 50 + index, 10 + index, 70 + index, 20 + index]
            )
    else:
        worksheet.append([None, datetime(year, 7, 28), None])
        worksheet.append(["LOCATION", "WET (g)", "DRY (g)"])
        for index, location in enumerate(LOCATIONS, start=1):
            worksheet.append([location, 100 + index, year - 1900 + index])


def _add_activity_sheet(
    workbook: Workbook,
    dates_by_year: dict[int, list[datetime | str | None]],
) -> None:
    worksheet = workbook.create_sheet("FIELD ACTIVITY DATES")
    years = sorted(dates_by_year)
    worksheet.append(["Field Activity Dates", *years])
    sample_count = max(len(dates) for dates in dates_by_year.values())
    for sample_number in range(1, sample_count + 1):
        worksheet.append(
            [
                f"Hay Sample {sample_number}",
                *[
                    dates_by_year[year][sample_number - 1]
                    if sample_number <= len(dates_by_year[year])
                    else None
                    for year in years
                ],
            ]
        )


class LabEtlBuilderTests(unittest.TestCase):
    def test_biomass_preserves_history_and_adds_only_new_dates(self) -> None:
        with tempfile.TemporaryDirectory() as directory_name:
            directory = Path(directory_name)
            workbook_path = directory / "master.xlsx"
            historical_path = directory / "historical.csv"
            workbook = Workbook()
            workbook.remove(workbook.active)
            _add_activity_sheet(
                workbook,
                {
                    2023: [datetime(2023, 6, 12), datetime(2023, 7, 25)],
                    2024: [datetime(2024, 7, 28)],
                    2025: [datetime(2025, 7, 28)],
                    2026: [datetime(2026, 7, 28)],
                },
            )
            _add_sheet(workbook, "2023 BIOMASS", 2023, old_layout=True)
            _add_sheet(workbook, "2024 BIOMASS", 2024)
            _add_sheet(workbook, "2025 BIOMASS", 2025)
            _add_sheet(workbook, "2026 BIOMASS", 2026)
            workbook.save(workbook_path)
            workbook.close()

            historical_path.write_text(
                "unnamed: 0,6/12/23,7/28/25\n"
                "Location,dry (g),dry (g)\n"
                + "\n".join(
                    f"{location},{index},{100 + index}"
                    for index, location in enumerate(LOCATIONS, 1)
                )
                + "\n",
                encoding="utf-8",
            )

            result = build_field_biomass(
                workbook_path=workbook_path,
                historical_path=historical_path,
            )

            self.assertEqual(
                result.columns.tolist(),
                ["location", "2023-06-12", "2025-07-28", "2026-07-28"],
            )
            self.assertEqual(result["location"].tolist(), LOCATIONS)
            self.assertEqual(result.loc[0, "2023-06-12"], 1)
            self.assertEqual(result.loc[0, "2026-07-28"], 127)

    def test_field_activity_date_replaces_stale_biomass_heading(self) -> None:
        with tempfile.TemporaryDirectory() as directory_name:
            workbook_path = Path(directory_name) / "master.xlsx"
            workbook = Workbook()
            workbook.remove(workbook.active)
            _add_activity_sheet(
                workbook,
                {
                    2023: [datetime(2023, 6, 12), datetime(2023, 7, 26)],
                },
            )
            _add_sheet(workbook, "2023 BIOMASS", 2023, old_layout=True)
            worksheet = workbook["2023 BIOMASS"]
            worksheet["G1"] = "9/6/23"
            worksheet["G2"] = "Location"
            worksheet["H2"] = "DRY"
            for row_number, location in enumerate(LOCATIONS, start=3):
                worksheet.cell(row=row_number, column=7, value=location)
                worksheet.cell(row=row_number, column=8, value=999)
            workbook.save(workbook_path)
            workbook.close()

            with self.assertWarnsRegex(
                BiomassDateWarning,
                "2023-07-25.*2023-07-26",
            ):
                result = extract_biomass_sheet(workbook_path, "2023 BIOMASS")

            self.assertEqual(
                sorted(result["sampling_date"].unique().tolist()),
                ["2023-07-26"],
            )

    def test_completed_biomass_event_requires_resolved_activity_date(self) -> None:
        with tempfile.TemporaryDirectory() as directory_name:
            workbook_path = Path(directory_name) / "master.xlsx"
            workbook = Workbook()
            workbook.remove(workbook.active)
            _add_activity_sheet(
                workbook,
                {2026: [datetime(2026, 7, 28), "TBD"]},
            )
            worksheet = workbook.create_sheet("2026 BIOMASS")
            worksheet.append(
                [
                    None,
                    datetime(2026, 7, 28),
                    None,
                    datetime(2026, 10, 16),
                    None,
                ]
            )
            worksheet.append(["LOCATION", "WET (g)", "DRY (g)", "WET (g)", "DRY (g)"])
            for index, location in enumerate(LOCATIONS, start=1):
                worksheet.append(
                    [location, 100 + index, 50 + index, 80 + index, 40 + index]
                )
            workbook.save(workbook_path)
            workbook.close()

            with self.assertRaisesRegex(ValueError, "refusing to publish"):
                extract_biomass_sheet(workbook_path, "2026 BIOMASS")

    def test_2026_nir_file_uses_filename_sampling_date(self) -> None:
        header_map = {
            "sample_id": "Sample ID 1",
            "crude_protein_pct_db": "Crude Protein Dry Basis",
        }
        with tempfile.TemporaryDirectory() as directory_name:
            source = Path(directory_name) / "NIR_2026-07-28.csv"
            pd.DataFrame(
                {
                    "Sample ID 1": ["S1HAY", "S2HAY", "S3HAY", "S4HAY"],
                    "Crude Protein Dry Basis": [10.6, 10.0, 12.6, 13.0],
                }
            ).to_csv(source, index=False)

            result = _read_supplemental_nir_csv(source, header_map)

            self.assertEqual(
                result["strip"].tolist(),
                ["strip_1", "strip_2", "strip_3", "strip_4"],
            )
            self.assertEqual(result["nir_date"].unique().tolist(), ["2026-07-28"])
            self.assertTrue(pd.to_numeric(result["crude_protein_pct_db"]).notna().all())

    def test_ward_results_filename_uses_compact_sampling_date(self) -> None:
        header_map = {
            "sample_id": "Sample ID 1",
            "crude_protein_pct_db": "Crude Protein Dry Basis",
        }
        with tempfile.TemporaryDirectory() as directory_name:
            source = Path(directory_name) / "NIR-Results-20261016.csv"
            pd.DataFrame(
                {
                    "Sample ID 1": ["S1HAY", "S2HAY", "S3HAY", "S4HAY"],
                    "Crude Protein Dry Basis": [15.9, 16.1, 15.4, 17.0],
                    "Date Received": ["10/22/2026"] * 4,
                }
            ).to_csv(source, index=False)

            result = _read_supplemental_nir_csv(source, header_map)

            self.assertEqual(result["nir_date"].unique().tolist(), ["2026-10-16"])
            self.assertEqual(
                result["strip"].tolist(),
                ["strip_1", "strip_2", "strip_3", "strip_4"],
            )

    def test_nir_table_includes_latest_year_from_clean_master(self) -> None:
        payload = build_nir_set1_table(WARD_MASTER_NIR_CSV)

        period_keys = [period["key"] for period in payload["periods"]]
        self.assertIn("2026-07-28", period_keys)
        self.assertIn("2026-10-16", period_keys)
        self.assertEqual(
            payload["data"]["crude_protein_pct_db"]["STRIP 1"]["2026-07-28"],
            10.6,
        )
        self.assertEqual(
            payload["data"]["crude_protein_pct_db"]["STRIP 1"]["2026-10-16"],
            15.9,
        )


if __name__ == "__main__":
    unittest.main()
